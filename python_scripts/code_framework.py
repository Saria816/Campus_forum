import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 创建一个新的Document对象
doc = docx.Document()

# 设置标题
title = doc.add_heading('2. 校园论坛代码基本框架', level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 添加正文
doc.add_heading('2.1 项目概述', level=2)
p = doc.add_paragraph('校园论坛是一个基于PHP开发的Web应用程序，采用MVC架构设计，为校园用户提供在线交流、讨论和互动的平台。论坛支持发帖、回复、私信、通知、分类浏览以及匿名发布等功能，并集成了AI助手功能，为用户提供智能化服务。')

doc.add_heading('2.2 技术架构', level=2)
p = doc.add_paragraph('本项目采用传统的LAMP(Linux + Apache + MySQL + PHP)技术栈开发，主要技术组件如下：')

# 添加技术架构列表
tech_stack = [
    ('前端技术', 'HTML5、CSS3、JavaScript'),
    ('CSS框架', 'Bootstrap 5.1.3 - 提供响应式布局和UI组件'),
    ('图标库', 'Font Awesome 6.1.1 - 提供丰富的图标资源'),
    ('JavaScript库', '原生JavaScript，辅以部分Bootstrap JS组件功能'),
    ('后端语言', 'PHP - 作为主要的服务器端脚本语言'),
    ('数据库', 'MySQL - 存储用户数据、帖子、评论等信息'),
    ('服务器环境', 'Apache/Nginx上的PHP环境')
]

table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '组件类型'
hdr_cells[1].text = '技术选择'

for tech, desc in tech_stack:
    row_cells = table.add_row().cells
    row_cells[0].text = tech
    row_cells[1].text = desc

doc.add_paragraph('')

doc.add_heading('2.3 系统架构', level=2)
p = doc.add_paragraph('系统采用经典的三层架构设计：')

architecture = [
    ('表示层（前端）', '负责用户界面展示，包括HTML模板、CSS样式和JavaScript交互'),
    ('业务逻辑层（后端）', '处理应用程序的核心业务逻辑，实现功能模块'),
    ('数据访问层', '通过PDO与MySQL数据库交互，实现数据的CRUD操作')
]

table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '架构层级'
hdr_cells[1].text = '职责描述'

for arch, desc in architecture:
    row_cells = table.add_row().cells
    row_cells[0].text = arch
    row_cells[1].text = desc

doc.add_paragraph('')

doc.add_heading('2.4 目录结构', level=2)
p = doc.add_paragraph('项目的文件组织结构如下：')

file_structure = [
    ('根目录', '存放主要的PHP页面文件，如index.php、forum.php等'),
    ('includes/', '包含可重用的PHP组件，如header.php、footer.php、functions.php和db.php等'),
    ('assets/css/', '存放样式表文件'),
    ('assets/js/', '存放JavaScript文件'),
    ('assets/images/', '存放网站图片资源')
]

table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '目录'
hdr_cells[1].text = '用途'

for dir, desc in file_structure:
    row_cells = table.add_row().cells
    row_cells[0].text = dir
    row_cells[1].text = desc

doc.add_paragraph('')

doc.add_heading('2.5 核心模块', level=2)
p = doc.add_paragraph('系统包含以下核心功能模块：')

core_modules = [
    ('用户管理', '用户注册、登录、个人资料管理'),
    ('帖子管理', '发布、编辑、删除和浏览帖子'),
    ('评论系统', '对帖子进行评论和回复'),
    ('分类系统', '按不同主题分类组织帖子'),
    ('消息系统', '用户之间的私信功能'),
    ('通知系统', '系统通知和用户互动通知'),
    ('匿名发布', '支持匿名发布帖子，保护用户隐私'),
    ('AI助手', '为用户提供AI对话支持'),
    ('管理后台', '管理员对用户、帖子、评论的管理功能'),
    ('敏感词过滤', '自动过滤不适当内容')
]

table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '模块名称'
hdr_cells[1].text = '功能描述'

for module, desc in core_modules:
    row_cells = table.add_row().cells
    row_cells[0].text = module
    row_cells[1].text = desc

doc.add_paragraph('')

doc.add_heading('2.6 数据库设计', level=2)
p = doc.add_paragraph('系统使用MySQL数据库，主要包含以下表：')

db_tables = [
    ('users', '存储用户信息，包括用户名、密码、邮箱等'),
    ('categories', '存储论坛分类信息'),
    ('posts', '存储帖子内容，包括标题、内容、作者、分类等'),
    ('comments', '存储帖子的评论信息'),
    ('post_likes', '记录用户对帖子的点赞'),
    ('messages', '存储用户之间的私信'),
    ('notifications', '存储系统和用户互动通知'),
    ('reports', '存储用户举报信息'),
    ('staff_messages', '存储用户向管理员发送的消息'),
    ('sensitive_words', '存储系统敏感词库')
]

table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = '表名'
hdr_cells[1].text = '用途'

for table_name, desc in db_tables:
    row_cells = table.add_row().cells
    row_cells[0].text = table_name
    row_cells[1].text = desc

doc.add_paragraph('')

doc.add_heading('2.7 安全措施', level=2)
p = doc.add_paragraph('系统实现了多种安全措施：')
security_measures = doc.add_paragraph()
security_measures.add_run('• PDO参数化查询：').bold = True
security_measures.add_run('防止SQL注入攻击\n')
security_measures.add_run('• 密码哈希存储：').bold = True
security_measures.add_run('使用安全的哈希算法存储用户密码\n')
security_measures.add_run('• XSS防护：').bold = True
security_measures.add_run('通过html转义函数(h())防止跨站脚本攻击\n')
security_measures.add_run('• CSRF防护：').bold = True
security_measures.add_run('表单提交时验证来源\n')
security_measures.add_run('• 会话管理：').bold = True
security_measures.add_run('安全的会话处理和用户认证')

doc.add_heading('2.8 辅助功能', level=2)
p = doc.add_paragraph('系统还包含一些辅助功能：')
auxiliary_features = doc.add_paragraph()
auxiliary_features.add_run('• 敏感词过滤：').bold = True
auxiliary_features.add_run('自动过滤不适当内容\n')
auxiliary_features.add_run('• 响应式设计：').bold = True
auxiliary_features.add_run('适配不同屏幕尺寸的设备\n')
auxiliary_features.add_run('• 邮箱域名验证：').bold = True
auxiliary_features.add_run('确保用户使用学校邮箱注册\n')
auxiliary_features.add_run('• 访问控制：').bold = True
auxiliary_features.add_run('基于角色的权限控制系统')

# 保存文档
doc.save('论坛代码基本框架.docx')
print("文档已保存为：论坛代码基本框架.docx") 